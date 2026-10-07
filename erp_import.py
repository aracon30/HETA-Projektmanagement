"""
Liest die 'Positionsansicht Auftrag.xlsm' (Auftrags-Dashboard, Netzlaufwerk
K:\\Datenstruktur\\Systemdateien\\DB_Statistik\\) ein und fasst die dort
positionsweise geführten Zeilen zu einem Kandidaten pro Auftrag (AB-Nummer)
zusammen, damit sie als Aufträge importiert werden können.

Diese Datei ist im Gegensatz zu MASTERDATEN AUFTRAG.xlsx / Auftragsübersicht.xlsm
nicht kennwortgeschützt (sie ist eine schreibgeschützte Momentaufnahme), daher
liest der Import sie direkt statt über die kennwortgeschützten Masterdateien.

Kommissionsnummer wird aus AB-Nummer + Jahr gebildet (K-<AB fünfstellig>/<Jahr
zweistellig>, z.B. AB 4869 im Jahr 2026 -> K-04869/26) — die Positionsansicht
enthält die Kommissionsnummer selbst nicht.
"""
from datetime import datetime, date

from openpyxl import load_workbook

REQUIRED_HEADERS = ["AB", "Jahr", "Kunde"]
KNOWN_HEADERS = REQUIRED_HEADERS + [
    "Position", "Positionsbeschreibung", "LT HETA AB", "Status", "Ordner",
    "Pfad Ordner", "Ursprungsauftrag", "zul. geliefert", "BU", "T",
]


def _cell_text(cell):
    if cell is None or cell.value is None:
        return ""
    return str(cell.value).strip()


def _parse_excel_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def hat_offene_position(kandidat):
    """True, wenn mindestens eine Position des Auftrags den Status 'offen' hat.
    Excel blendet abgeschlossene/historische Positionen per Filter aus, aber
    openpyxl liest ausgefilterte Zeilen trotzdem mit - dieser Check ersetzt
    den Excel-Filter beim Import, sonst kommen jahrelange Altaufträge mit."""
    return any(s.strip().lower() == "offen" for s in kandidat["statusWerte"])


def _find_header_row(sheet):
    for row in sheet.iter_rows(min_row=1, max_row=30):
        header_map = {}
        for cell in row:
            text = _cell_text(cell)
            if text in KNOWN_HEADERS:
                header_map[text] = cell.column
        if "AB" in header_map and "Kunde" in header_map:
            return row[0].row, header_map
    return None, None


def parse_positionsansicht_auftrag(file):
    """file: Dateipfad oder file-artiges Objekt (z.B. Flask FileStorage.stream)."""
    wb = load_workbook(file, data_only=True)

    sheet = None
    for name in wb.sheetnames:
        if "position" in name.lower():
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]

    header_row, header_map = _find_header_row(sheet)
    if header_row is None:
        raise ValueError(
            "Kopfzeile mit Spalten 'AB' und 'Kunde' nicht gefunden — ist das die "
            "Datei 'Positionsansicht Auftrag.xlsm' (Blatt 'Positionen Auftrag')?"
        )
    missing = [h for h in REQUIRED_HEADERS if h not in header_map]
    if missing:
        raise ValueError(f"Erwartete Spalte(n) nicht gefunden: {', '.join(missing)}")

    groups = {}
    order = []

    for row in sheet.iter_rows(min_row=header_row + 1):
        def cell(name):
            col = header_map.get(name)
            return row[col - 1] if col else None

        ab_cell = cell("AB")
        ab_text = _cell_text(ab_cell)
        if not ab_text:
            continue
        try:
            ab_number = int(ab_text)
        except ValueError:
            continue

        if ab_number not in groups:
            order.append(ab_number)
            groups[ab_number] = {
                "erpAbNummer": str(ab_number),
                "jahr": None,
                "kunde": "",
                "liefertermin": None,
                "ordnerPfad": None,
                "positionenText": [],
                "positionenDetail": [],
                "statusWerte": [],
                "ursprungsauftrag": "",
                "zulGeliefert": "",
                "bu": "",
                "t": "",
            }
        g = groups[ab_number]

        jahr_text = _cell_text(cell("Jahr"))
        if jahr_text and g["jahr"] is None:
            try:
                g["jahr"] = int(jahr_text)
            except ValueError:
                pass

        kunde_text = _cell_text(cell("Kunde"))
        if kunde_text and not g["kunde"]:
            g["kunde"] = kunde_text

        pos_text = _cell_text(cell("Position"))
        beschr_text = _cell_text(cell("Positionsbeschreibung"))
        lt_cell = cell("LT HETA AB")
        pos_liefertermin = (
            _parse_excel_date(lt_cell.value)
            if lt_cell is not None and lt_cell.value not in (None, "")
            else None
        )
        if beschr_text:
            g["positionenText"].append(f"{pos_text}: {beschr_text}" if pos_text else beschr_text)
            g["positionenDetail"].append({
                "position": pos_text or None,
                "beschreibung": beschr_text,
                "liefertermin": pos_liefertermin,
            })
            # Liefertermin auf Auftragsebene = frühester Termin unter den Positionen
            # (Positionen eines Auftrags können unterschiedliche Liefertermine haben,
            # z.B. bei Teillieferungen — nicht einfach den ersten gefundenen nehmen).
            if pos_liefertermin and (g["liefertermin"] is None or pos_liefertermin < g["liefertermin"]):
                g["liefertermin"] = pos_liefertermin

        if not g["ordnerPfad"]:
            pfad_text = _cell_text(cell("Pfad Ordner"))
            if pfad_text:
                g["ordnerPfad"] = pfad_text
            else:
                # Fallback für Dateien, in denen "Ordner" noch ein echter
                # Excel-Hyperlink ist statt einer HYPERLINK()-Formel, die nur
                # den Anzeigetext liefert (dann steht der Pfad in "Pfad Ordner").
                ordner_cell = cell("Ordner")
                if ordner_cell is not None and ordner_cell.hyperlink:
                    g["ordnerPfad"] = ordner_cell.hyperlink.target

        status_text = _cell_text(cell("Status"))
        if status_text and status_text not in g["statusWerte"]:
            g["statusWerte"].append(status_text)

        ursprung_text = _cell_text(cell("Ursprungsauftrag"))
        if ursprung_text and not g["ursprungsauftrag"]:
            g["ursprungsauftrag"] = ursprung_text

        zul_text = _cell_text(cell("zul. geliefert"))
        if zul_text and not g["zulGeliefert"]:
            g["zulGeliefert"] = zul_text

        bu_text = _cell_text(cell("BU"))
        if bu_text and not g["bu"]:
            g["bu"] = bu_text

        t_text = _cell_text(cell("T"))
        if t_text and not g["t"]:
            g["t"] = t_text

    results = []
    for ab_number in order:
        g = groups[ab_number]
        jahr = g["jahr"] or datetime.now().year
        kommission = f"K-{ab_number:05d}/{jahr % 100:02d}"
        results.append({
            "erpAbNummer": g["erpAbNummer"],
            "abNummer": ab_number,
            "jahr": jahr,
            "kommission": kommission,
            "kunde": g["kunde"],
            "lieferumfang": "\n".join(g["positionenText"]),
            "liefertermin": g["liefertermin"],
            "ordnerPfad": g["ordnerPfad"],
            "positionsAnzahl": len(g["positionenText"]),
            "positionen": [
                {
                    "position": p["position"],
                    "beschreibung": p["beschreibung"],
                    "liefertermin": p["liefertermin"].isoformat() if p["liefertermin"] else None,
                }
                for p in g["positionenDetail"]
            ],
            "statusWerte": g["statusWerte"],
            "ursprungsauftrag": g["ursprungsauftrag"] or None,
            "zulGeliefert": g["zulGeliefert"] or None,
            "bu": g["bu"] or None,
            "t": g["t"] or None,
        })
    return results
